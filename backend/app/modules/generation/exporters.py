import os
import io
import re
import docx
from abc import ABC, abstractmethod
from typing import Dict, Any, Tuple
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


class BaseExporter(ABC):
    @abstractmethod
    def export(self, cdm_data: Dict[str, Any], publisher_code: str = "IEEE") -> Tuple[bytes, str, str]:
        """
        Exports CDM to file bytes.
        Returns: (file_bytes, file_extension, mime_type)
        """
        pass


class PdfExporter(BaseExporter):
    def export(self, cdm_data: Dict[str, Any], publisher_code: str = "IEEE") -> Tuple[bytes, str, str]:
        metadata = cdm_data.get("metadata", {})
        sections = cdm_data.get("sections", [])
        references = cdm_data.get("references", [])
        keywords = metadata.get("keywords", [])
        pub = (publisher_code or "IEEE").upper()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=45,
            leftMargin=45,
            topMargin=45,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18 if pub == "IEEE" else 16,
            leading=22,
            alignment=1,  # Centered
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=10
        )
        author_style = ParagraphStyle(
            'DocAuthor',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=13,
            alignment=1,
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=4
        )
        affil_style = ParagraphStyle(
            'DocAffil',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=8.5,
            leading=11,
            alignment=1,
            textColor=colors.HexColor('#475569'),
            spaceAfter=12
        )
        heading_style = ParagraphStyle(
            'DocHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11 if pub == "IEEE" else 12,
            leading=15,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=14,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=9.5 if pub == "IEEE" else 10,
            leading=13.5,
            alignment=4,  # Justified
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=8
        )
        abstract_style = ParagraphStyle(
            'DocAbstract',
            parent=body_style,
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            leftIndent=15,
            rightIndent=15,
            spaceAfter=6
        )

        story = []

        # Publisher Header Badge / Banner
        if pub == "IEEE":
            story.append(Paragraph("<b>IEEE CONFERENCE / TRANSACTIONS RESEARCH PUBLICATION</b>", ParagraphStyle('PubTag', fontName='Helvetica-Bold', fontSize=8, textColor=colors.HexColor('#002855'), alignment=1)))
        else:
            story.append(Paragraph(f"<b>[{pub} Format Export]</b>", ParagraphStyle('PubTag', fontName='Helvetica-Bold', fontSize=8, textColor=colors.HexColor('#2563EB'), alignment=1)))
        story.append(Spacer(1, 8))

        # Title
        story.append(Paragraph(metadata.get("title", "Research Paper"), title_style))

        # Authors & Affiliations
        authors = metadata.get("authors", [])
        if authors:
            auth_names = ", ".join([a.get("name", "Author") for a in authors])
            story.append(Paragraph(auth_names, author_style))
            affil_strs = []
            for a in authors:
                part = a.get("affiliation", "")
                if a.get("email"):
                    part += f" ({a.get('email')})"
                if part and part not in affil_strs:
                    affil_strs.append(part)
            if affil_strs:
                story.append(Paragraph(" | ".join(affil_strs), affil_style))
        story.append(Spacer(1, 4))
        story.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor('#CBD5E1'), spaceAfter=10))

        # Abstract (with IEEE Abstract em-dash)
        abstract = metadata.get("abstract", "")
        if abstract:
            if pub == "IEEE":
                story.append(Paragraph(f"<b><i>Abstract</i>—</b><i>{abstract}</i>", abstract_style))
            else:
                story.append(Paragraph("<b>Abstract</b>", heading_style))
                story.append(Paragraph(f"<i>{abstract}</i>", abstract_style))

        # Keywords / Index Terms
        if keywords:
            kw_str = ", ".join(keywords)
            if pub == "IEEE":
                story.append(Paragraph(f"<b><i>Index Terms</i>—</b><i>{kw_str}</i>", abstract_style))
            else:
                story.append(Paragraph(f"<b>Keywords:</b> {kw_str}", abstract_style))
            story.append(Spacer(1, 8))

        # Sections (ensuring Roman numeral numbering for IEEE)
        roman_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
        for idx, sec in enumerate(sections):
            sec_title = sec.get("title", "Section").strip()
            if pub == "IEEE" and not any(sec_title.startswith(f"{r}.") for r in roman_numerals):
                r_num = roman_numerals[idx] if idx < len(roman_numerals) else f"{idx+1}"
                sec_title = f"{r_num}. {sec_title.upper()}"
            story.append(Paragraph(sec_title, heading_style))
            content = sec.get("content", "")
            for p in content.split("\n\n"):
                if p.strip():
                    story.append(Paragraph(p.strip(), body_style))

        # Tables
        tables = cdm_data.get("tables", [])
        if tables:
            story.append(Spacer(1, 10))
            for idx, tbl in enumerate(tables, 1):
                caption = tbl.get("caption") or f"Table {idx}"
                if pub == "IEEE":
                    caption = f"TABLE {roman_numerals[idx-1] if idx-1 < len(roman_numerals) else idx}: {caption.upper()}"
                story.append(Paragraph(f"<b>{caption}</b>", heading_style))
                headers = tbl.get("headers", [])
                rows = tbl.get("rows", [])
                if not rows and tbl.get("data"):
                    raw_data = tbl.get("data", [])
                    if raw_data:
                        headers = raw_data[0]
                        rows = raw_data[1:]

                table_matrix = []
                if headers:
                    table_matrix.append([Paragraph(f"<b>{str(h)}</b>", body_style) for h in headers])
                for r in rows:
                    if isinstance(r, list):
                        table_matrix.append([Paragraph(str(c), body_style) for c in r])
                if table_matrix:
                    rl_tbl = Table(table_matrix, hAlign='CENTER')
                    rl_tbl.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F8FAFC')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                        ('TOPPADDING', (0, 0), (-1, -1), 4),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                    ]))
                    story.append(rl_tbl)
                    story.append(Spacer(1, 8))

        # References
        if references:
            story.append(Spacer(1, 10))
            ref_heading = "REFERENCES" if pub == "IEEE" else "References"
            story.append(Paragraph(f"<b>{ref_heading}</b>", heading_style))
            for r_idx, ref in enumerate(references, start=1):
                key = ref.get("key") or f"[{r_idx}]"
                ref_text = ref.get("text") or f"{key} {ref.get('title', 'Reference item')}, {ref.get('venue', '')}, {ref.get('year', '')}."
                if not ref_text.startswith(key):
                    ref_text = f"{key} {ref_text}"
                story.append(Paragraph(ref_text, ParagraphStyle('RefStyle', parent=body_style, fontSize=8.5, leading=11, spaceAfter=4)))

        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes, ".pdf", "application/pdf"


class DocxExporter(BaseExporter):
    def export(self, cdm_data: Dict[str, Any], publisher_code: str = "IEEE") -> Tuple[bytes, str, str]:
        metadata = cdm_data.get("metadata", {})
        sections = cdm_data.get("sections", [])
        references = cdm_data.get("references", [])
        keywords = metadata.get("keywords", [])
        pub = (publisher_code or "IEEE").upper()

        doc = docx.Document()
        roman_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]

        # Header / Title
        p_title = doc.add_heading(metadata.get("title", "Research Paper"), level=0)
        p_title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER

        # Authors
        authors = metadata.get("authors", [])
        if authors:
            auth_str = ", ".join([a.get("name", "Author") for a in authors])
            p_auth = doc.add_paragraph(auth_str)
            p_auth.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            affils = [a.get("affiliation", "") for a in authors if a.get("affiliation")]
            if affils:
                p_affil = doc.add_paragraph(" | ".join(dict.fromkeys(affils)))
                p_affil.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                p_affil.runs[0].italic = True

        # Abstract
        abstract = metadata.get("abstract", "")
        if abstract:
            p_abs = doc.add_paragraph()
            r_label = p_abs.add_run("Abstract—" if pub == "IEEE" else "Abstract: ")
            r_label.bold = True
            r_label.italic = True
            r_text = p_abs.add_run(abstract)
            r_text.italic = True

        # Index Terms / Keywords
        if keywords:
            p_kw = doc.add_paragraph()
            r_kw_lbl = p_kw.add_run("Index Terms—" if pub == "IEEE" else "Keywords: ")
            r_kw_lbl.bold = True
            r_kw_lbl.italic = True
            p_kw.add_run(", ".join(keywords))

        # Sections
        for idx, sec in enumerate(sections):
            sec_title = sec.get("title", "Section").strip()
            if pub == "IEEE" and not any(sec_title.startswith(f"{r}.") for r in roman_numerals):
                r_num = roman_numerals[idx] if idx < len(roman_numerals) else f"{idx+1}"
                sec_title = f"{r_num}. {sec_title.upper()}"
            doc.add_heading(sec_title, level=sec.get("level", 1))
            content = sec.get("content", "")
            for para in content.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())

        # Tables
        tables = cdm_data.get("tables", [])
        for idx, tbl in enumerate(tables, 1):
            caption = tbl.get("caption") or f"Table {idx}"
            if pub == "IEEE":
                caption = f"TABLE {roman_numerals[idx-1] if idx-1 < len(roman_numerals) else idx}: {caption.upper()}"
            doc.add_heading(caption, level=2)
            headers = tbl.get("headers", [])
            rows = tbl.get("rows", [])
            if not rows and tbl.get("data"):
                raw_data = tbl.get("data", [])
                if raw_data:
                    headers = raw_data[0]
                    rows = raw_data[1:]

            num_cols = max(len(headers), max((len(r) for r in rows), default=0))
            if num_cols > 0:
                t = doc.add_table(rows=0, cols=num_cols)
                t.style = 'Table Grid'
                if headers:
                    hdr_cells = t.add_row().cells
                    for i, h in enumerate(headers):
                        if i < num_cols:
                            hdr_cells[i].text = str(h)
                for r in rows:
                    if isinstance(r, list):
                        row_cells = t.add_row().cells
                        for i, c in enumerate(r):
                            if i < num_cols:
                                row_cells[i].text = str(c)

        # References
        if references:
            ref_heading = "REFERENCES" if pub == "IEEE" else "References"
            doc.add_heading(ref_heading, level=1)
            for r_idx, ref in enumerate(references, start=1):
                key = ref.get("key") or f"[{r_idx}]"
                ref_text = ref.get("text") or f"{key} {ref.get('title', 'Reference item')}, {ref.get('venue', '')}, {ref.get('year', '')}."
                if not ref_text.startswith(key):
                    ref_text = f"{key} {ref_text}"
                doc.add_paragraph(ref_text)

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes, ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class LatexExporter(BaseExporter):
    def export(self, cdm_data: Dict[str, Any], publisher_code: str = "IEEE") -> Tuple[bytes, str, str]:
        metadata = cdm_data.get("metadata", {})
        sections = cdm_data.get("sections", [])
        references = cdm_data.get("references", [])
        tables = cdm_data.get("tables", [])
        keywords = metadata.get("keywords", [])

        title = metadata.get("title", "Research Document")
        abstract = metadata.get("abstract", "")
        authors = metadata.get("authors", [])

        pub = publisher_code.upper() if publisher_code else "IEEE"
        auth_latex = " \\and ".join([a.get("name", "Author") for a in authors]) if authors else "Author"

        if pub == "ACM":
            doc_class = "\\documentclass[sigconf]{acmart}"
            author_block = "\n".join([f"\\author{{{a.get('name', 'Author')}}}\n\\affiliation{{\\institution{{{a.get('affiliation', 'Institution')}}}}}" for a in authors]) if authors else "\\author{Author}"
        elif pub == "ELSEVIER":
            doc_class = "\\documentclass[preprint,12pt]{elsarticle}"
            author_block = f"\\author{{{auth_latex}}}"
        elif pub == "SPRINGER":
            doc_class = "\\documentclass{llncs}"
            author_block = f"\\author{{{auth_latex}}}"
        elif pub == "NATURE":
            doc_class = "\\documentclass[11pt]{article}\n\\usepackage[margin=1in]{geometry}"
            author_block = f"\\author{{{auth_latex}}}"
        else:  # IEEE default
            doc_class = "\\documentclass[conference]{IEEEtran}"
            if authors:
                auth_blocks = []
                for a in authors:
                    a_name = a.get("name", "Author Name")
                    a_affil = a.get("affiliation", "Department of Engineering \\\\ Institution")
                    a_email = a.get("email", "")
                    block = f"\\IEEEauthorblockN{{{a_name}}}\n\\IEEEauthorblockA{{\\textit{{{a_affil}}}"
                    if a_email:
                        block += f"\\\\\n{a_email}"
                    block += "}"
                    auth_blocks.append(block)
                author_block = "\\author{" + " \\and\n".join(auth_blocks) + "}"
            else:
                author_block = "\\author{\\IEEEauthorblockN{Author Name}\\IEEEauthorblockA{\\textit{Department of Engineering}\\\\Research Institution}}"

        tex_lines = [
            doc_class,
            "\\usepackage{cite}",
            "\\usepackage{amsmath,amssymb,amsfonts}",
            "\\usepackage{graphicx}",
            "\\usepackage{booktabs}",
            "\\begin{document}",
            f"\\title{{{title}}}",
            author_block,
            "\\maketitle",
        ]

        if abstract:
            tex_lines.extend([
                "\\begin{abstract}",
                abstract,
                "\\end{abstract}"
            ])

        if keywords:
            kw_str = ", ".join(keywords)
            if pub == "IEEE":
                tex_lines.extend([
                    "\\begin{IEEEkeywords}",
                    kw_str,
                    "\\end{IEEEkeywords}"
                ])

        roman_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
        for idx, sec in enumerate(sections):
            sec_title = sec.get("title", "Section").strip()
            # If IEEE, clean pre-existing roman numeral for LaTeX \section as IEEEtran handles section numbering
            if pub == "IEEE":
                sec_title = re.sub(r"^(I|II|III|IV|V|VI|VII|VIII|IX|X|\d+)\.?\s*", "", sec_title).strip()
            tex_lines.append(f"\\section{{{sec_title}}}")
            content = sec.get("content", "")
            tex_lines.append(content)

        # Render Tables if any
        for tbl in tables:
            caption = tbl.get("caption", "Table")
            headers = tbl.get("headers", [])
            rows = tbl.get("rows", [])
            if not rows and tbl.get("data"):
                raw_data = tbl.get("data", [])
                if raw_data:
                    headers = raw_data[0]
                    rows = raw_data[1:]

            if headers or rows:
                cols_count = max(len(headers), max((len(r) for r in rows), default=0))
                col_spec = "c" * cols_count
                tbl_tex = [
                    "\\begin{table}[htbp]",
                    f"\\caption{{{caption.upper() if pub == 'IEEE' else caption}}}",
                    "\\centering",
                    f"\\begin{{tabular}}{{{col_spec}}}",
                    "\\toprule"
                ]
                if headers:
                    tbl_tex.append(" & ".join(f"\\textbf{{{h}}}" for h in headers) + " \\\\")
                    tbl_tex.append("\\midrule")
                for row in rows:
                    if isinstance(row, list):
                        tbl_tex.append(" & ".join(str(cell) for cell in row) + " \\\\")
                tbl_tex.append("\\bottomrule")
                tbl_tex.append("\\end{tabular}")
                tbl_tex.append("\\end{table}")
                tex_lines.append("\n".join(tbl_tex))

        if references:
            tex_lines.extend([
                "\\begin{thebibliography}{00}",
            ])
            for idx, ref in enumerate(references, start=1):
                ref_text = ref.get("text", "Reference item")
                tex_lines.append(f"\\bibitem{{b{idx}}} {ref_text}")
            tex_lines.append("\\end{thebibliography}")

        tex_lines.append("\\end{document}")

        latex_str = "\n\n".join(tex_lines)
        return latex_str.encode("utf-8"), ".tex", "application/x-tex"


class HtmlExporter(BaseExporter):
    def export(self, cdm_data: Dict[str, Any], publisher_code: str = "IEEE") -> Tuple[bytes, str, str]:
        metadata = cdm_data.get("metadata", {})
        sections = cdm_data.get("sections", [])
        references = cdm_data.get("references", [])
        tables = cdm_data.get("tables", [])
        keywords = metadata.get("keywords", [])
        pub = (publisher_code or "IEEE").upper()

        title = metadata.get("title", "Research Document")
        abstract = metadata.get("abstract", "")
        authors = metadata.get("authors", [])
        authors_html = ", ".join([f"<span class='author-name'>{a.get('name', 'Author')}</span>" for a in authors])

        sections_html = ""
        roman_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
        for idx, sec in enumerate(sections):
            sec_title = sec.get("title", "Section").strip()
            if pub == "IEEE" and not any(sec_title.startswith(f"{r}.") for r in roman_numerals):
                r_num = roman_numerals[idx] if idx < len(roman_numerals) else f"{idx+1}"
                sec_title = f"{r_num}. {sec_title.upper()}"
            sections_html += f"<section><h2>{sec_title}</h2><p>{sec.get('content', '').replace(chr(10), '<br>')}</p></section>"

        tables_html = ""
        for idx, tbl in enumerate(tables, 1):
            caption = tbl.get("caption") or f"Table {idx}"
            if pub == "IEEE":
                caption = f"TABLE {roman_numerals[idx-1] if idx-1 < len(roman_numerals) else idx}: {caption.upper()}"
            headers = tbl.get("headers", [])
            rows = tbl.get("rows", [])
            if not rows and tbl.get("data"):
                raw_data = tbl.get("data", [])
                if raw_data:
                    headers = raw_data[0]
                    rows = raw_data[1:]

            header_html = "".join([f"<th>{h}</th>" for h in headers]) if headers else ""
            row_html = "".join(["<tr>" + "".join([f"<td>{cell}</td>" for cell in r]) + "</tr>" for r in rows if isinstance(r, list)])
            tables_html += f"<div class='table-container'><h3>{caption}</h3><table><thead><tr>{header_html}</tr></thead><tbody>{row_html}</tbody></table></div>"

        references_html = ""
        if references:
            ref_items = "".join([f"<li>{r.get('text', 'Reference')}</li>" for r in references])
            references_html = f"<section><h2>{'REFERENCES' if pub == 'IEEE' else 'References'}</h2><ol class='references-list'>{ref_items}</ol></section>"

        kw_html = f"<div class='index-terms'><strong><em>Index Terms</em>—</strong><em>{', '.join(keywords)}</em></div>" if keywords else ""

        html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>{title}</title>
    <style>
        body {{ font-family: 'Times New Roman', Times, serif; margin: 40px auto; max-width: 850px; color: #111; line-height: 1.5; padding: 0 20px; }}
        header {{ text-align: center; border-bottom: 1px solid #ddd; padding-bottom: 16px; margin-bottom: 20px; }}
        h1 {{ font-size: 24px; font-weight: bold; margin-bottom: 10px; }}
        .pub-badge {{ font-family: sans-serif; font-size: 11px; font-weight: bold; color: #002855; margin-bottom: 8px; text-transform: uppercase; }}
        .authors {{ font-size: 13px; font-weight: bold; color: #222; margin-bottom: 12px; }}
        .abstract {{ font-size: 12px; line-height: 1.6; margin-bottom: 16px; text-align: justify; padding: 0 10px; }}
        .index-terms {{ font-size: 12px; margin-bottom: 24px; padding: 0 10px; }}
        h2 {{ font-size: 15px; font-weight: bold; text-transform: uppercase; border-bottom: 1px solid #eee; padding-bottom: 3px; color: #0f172a; margin-top: 24px; }}
        p {{ font-size: 12.5px; text-align: justify; margin-bottom: 12px; }}
        table {{ border-collapse: collapse; width: 100%; margin: 16px 0; font-size: 12px; }}
        th {{ background: #f8fafc; border: 1px solid #cbd5e1; padding: 6px; text-align: center; font-weight: bold; }}
        td {{ border: 1px solid #cbd5e1; padding: 6px; text-align: center; }}
        .table-container h3 {{ font-size: 12px; text-align: center; margin-bottom: 6px; text-transform: uppercase; font-weight: bold; }}
        .references-list {{ font-size: 11px; padding-left: 20px; }}
        .references-list li {{ margin-bottom: 6px; text-align: justify; }}
    </style>
</head>
<body>
    <header>
        <div class="pub-badge">IEEE Research Publication Template</div>
        <h1>{title}</h1>
        <div class="authors">{authors_html}</div>
    </header>
    {f'<div class="abstract"><strong><em>Abstract</em>—</strong><em>{abstract}</em></div>' if abstract else ''}
    {kw_html}
    <main>
        {sections_html}
        {tables_html}
        {references_html}
    </main>
</body>
</html>"""

        return html_content.encode("utf-8"), ".html", "text/html"


class ExporterFactory:
    @staticmethod
    def get_exporter(export_format: str) -> BaseExporter:
        fmt = export_format.upper()
        if fmt == "PDF":
            return PdfExporter()
        elif fmt in ["DOCX", "DOC"]:
            return DocxExporter()
        elif fmt in ["LATEX", "TEX"]:
            return LatexExporter()
        elif fmt == "HTML":
            return HtmlExporter()
        else:
            return PdfExporter()
