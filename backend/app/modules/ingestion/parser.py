import os
import re
import fitz  # PyMuPDF
import docx
from typing import Dict, Any, List


class DocumentParser:
    """
    Parses PDF, DOCX, TXT, MD, and LaTeX files into a standard raw structure suitable for CDM generation.
    """

    @classmethod
    def parse_file(cls, file_path: str, file_type: str) -> Dict[str, Any]:
        ext = file_type.lower().lstrip(".")
        if ext == "pdf":
            return cls._parse_pdf(file_path)
        elif ext in ["docx", "doc"]:
            return cls._parse_docx(file_path)
        elif ext in ["txt", "text"]:
            return cls._parse_txt(file_path)
        elif ext in ["md", "markdown"]:
            return cls._parse_markdown(file_path)
        elif ext in ["tex", "latex"]:
            return cls._parse_latex(file_path)
        else:
            return cls._parse_txt(file_path)

    @classmethod
    def _parse_pdf(cls, file_path: str) -> Dict[str, Any]:
        doc = fitz.open(file_path)
        full_text_blocks = []
        pages_text = []

        for page in doc:
            p_text = page.get_text("text")
            pages_text.append(p_text)
            full_text_blocks.append(p_text)

        full_text = "\n\n".join(full_text_blocks)
        doc.close()

        lines = [line.strip() for line in full_text.split("\n") if line.strip()]
        title = lines[0] if lines else "Untitled Document"
        
        # Simple metadata extract
        abstract = ""
        abstract_match = re.search(r"(?i)abstract[:\s]*(.*?)(?=\n\s*\n|introduction|1\.|keywords)", full_text, re.DOTALL)
        if abstract_match:
            abstract = abstract_match.group(1).strip()

        keywords = []
        kw_match = re.search(r"(?i)keywords[:\s]*(.*?)(?=\n\s*\n|1\.|introduction)", full_text, re.DOTALL)
        if kw_match:
            raw_kws = kw_match.group(1).strip()
            keywords = [k.strip() for k in re.split(r"[,;•]", raw_kws) if k.strip()]

        sections = cls._structure_raw_text(full_text)

        return {
            "title": title,
            "abstract": abstract,
            "keywords": keywords,
            "authors": [{"name": "Extracted Author", "affiliation": "", "email": ""}],
            "raw_text": full_text,
            "sections": sections,
            "figures": [],
            "tables": []
        }

    @classmethod
    def _parse_docx(cls, file_path: str) -> Dict[str, Any]:
        doc = docx.Document(file_path)
        full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        
        title = "Untitled Document"
        for p in doc.paragraphs:
            if p.text.strip():
                title = p.text.strip()
                break

        abstract = ""
        keywords = []
        
        abstract_match = re.search(r"(?i)abstract[:\s]*(.*?)(?=\n\s*\n|introduction|1\.|keywords)", full_text, re.DOTALL)
        if abstract_match:
            abstract = abstract_match.group(1).strip()

        kw_match = re.search(r"(?i)keywords[:\s]*(.*?)(?=\n\s*\n|1\.|introduction)", full_text, re.DOTALL)
        if kw_match:
            raw_kws = kw_match.group(1).strip()
            keywords = [k.strip() for k in re.split(r"[,;•]", raw_kws) if k.strip()]

        sections = cls._structure_raw_text(full_text)

        # Extract tables
        tables = []
        for t_idx, tbl in enumerate(doc.tables, 1):
            table_data = []
            for row in tbl.rows:
                table_data.append([cell.text.strip() for cell in row.cells])
            if table_data:
                tables.append({
                    "id": f"tbl_{t_idx}",
                    "caption": f"Table {t_idx}",
                    "data": table_data
                })

        return {
            "title": title,
            "abstract": abstract,
            "keywords": keywords,
            "authors": [{"name": "Extracted Author", "affiliation": "", "email": ""}],
            "raw_text": full_text,
            "sections": sections,
            "tables": tables,
            "figures": []
        }

    @classmethod
    def _parse_txt(cls, file_path: str) -> Dict[str, Any]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read()

        lines = [line.strip() for line in full_text.split("\n") if line.strip()]
        title = lines[0] if lines else "Untitled Document"

        title_match = re.search(r"(?i)^title[:\s]+(.*)$", full_text, re.MULTILINE)
        if title_match:
            title = title_match.group(1).strip()

        abstract = ""
        abstract_match = re.search(r"(?i)abstract[:\s]*(.*?)(?=\n\s*\n|introduction|1\.|keywords)", full_text, re.DOTALL)
        if abstract_match:
            abstract = abstract_match.group(1).strip()

        keywords = []
        kw_match = re.search(r"(?i)keywords[:\s]*(.*?)(?=\n\s*\n|1\.|introduction)", full_text, re.DOTALL)
        if kw_match:
            raw_kws = kw_match.group(1).strip()
            keywords = [k.strip() for k in re.split(r"[,;•]", raw_kws) if k.strip()]

        authors = [{"name": "Lead Researcher", "affiliation": "Academic Institution", "email": ""}]
        author_match = re.search(r"(?i)authors?[:\s]*(.*?)(?=\n\s*\n|abstract|keywords|1\.|introduction)", full_text, re.DOTALL)
        if author_match:
            raw_auths = author_match.group(1).strip()
            ext_authors = []
            for a in re.split(r"[,;\n]", raw_auths):
                if a.strip() and len(a.strip()) > 2:
                    ext_authors.append({"name": a.strip(), "affiliation": "", "email": ""})
            if ext_authors:
                authors = ext_authors

        sections = cls._structure_raw_text(full_text)
        return {
            "title": title,
            "abstract": abstract,
            "keywords": keywords,
            "authors": authors,
            "raw_text": full_text,
            "sections": sections,
            "tables": [],
            "figures": []
        }

    @classmethod
    def _parse_markdown(cls, file_path: str) -> Dict[str, Any]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read()

        title = "Untitled Document"
        title_match = re.search(r"^#\s+(.*)$", full_text, re.MULTILINE)
        if title_match:
            title = title_match.group(1).strip()

        abstract = ""
        abstract_match = re.search(r"(?i)(?:##\s*abstract|abstract)[:\s]*(.*?)(?=\n#{1,3}\s+|\n\s*\n\s*#{1,3})", full_text, re.DOTALL)
        if abstract_match:
            abstract = abstract_match.group(1).strip()

        keywords = []
        kw_match = re.search(r"(?i)keywords[:\s]*(.*?)(?=\n\s*\n|\n#{1,3}\s+)", full_text, re.DOTALL)
        if kw_match:
            raw_kws = kw_match.group(1).strip()
            keywords = [k.strip() for k in re.split(r"[,;•`]", raw_kws) if k.strip()]

        sections = []
        section_splits = re.split(r"\n(?=#{1,3}\s+)", full_text)
        
        sec_id = 1
        for block in section_splits:
            if not block.strip():
                continue
            header_match = re.match(r"^(#{1,3})\s+(.*)$", block, re.MULTILINE)
            if header_match:
                level = len(header_match.group(1))
                sec_title = header_match.group(2).strip()
                content = block[header_match.end():].strip()
            else:
                level = 1
                sec_title = f"Section {sec_id}"
                content = block.strip()

            sections.append({
                "id": f"sec_{sec_id}",
                "title": sec_title,
                "level": level,
                "content": content,
                "order": sec_id,
                "subsections": []
            })
            sec_id += 1

        return {
            "title": title,
            "abstract": abstract,
            "keywords": keywords,
            "authors": [{"name": "Author", "affiliation": "", "email": ""}],
            "raw_text": full_text,
            "sections": sections,
            "tables": [],
            "figures": []
        }

    @classmethod
    def _parse_latex(cls, file_path: str) -> Dict[str, Any]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read()

        title_match = re.search(r"\\title\{([^}]+)\}", full_text)
        title = title_match.group(1).strip() if title_match else "Untitled LaTeX Document"

        abstract_match = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", full_text, re.DOTALL)
        abstract = abstract_match.group(1).strip() if abstract_match else ""

        # Authors
        authors = []
        author_match = re.search(r"\\author\{([^}]+)\}", full_text)
        if author_match:
            raw_authors = author_match.group(1).strip()
            for auth_part in re.split(r"\\and|,|\\\\", raw_authors):
                clean_name = re.sub(r"\\[a-zA-Z]+(\{[^}]*\})?", "", auth_part).strip()
                if clean_name:
                    authors.append({"name": clean_name, "affiliation": "", "email": ""})
        if not authors:
            authors = [{"name": "Lead Author", "affiliation": "Institution", "email": ""}]

        # Keywords
        keywords = []
        kw_match = re.search(r"(?:\\begin\{IEEEkeywords\}|\\keywords\{)(.*?)(?:\\end\{IEEEkeywords\}|\})", full_text, re.DOTALL)
        if kw_match:
            raw_kws = kw_match.group(1).strip()
            keywords = [k.strip() for k in re.split(r"[,;•\\]", raw_kws) if k.strip() and not k.startswith("\\")]

        # Extract sections with real content
        section_pattern = r"\\(section|subsection|subsubsection)\*?\{([^}]+)\}"
        matches = list(re.finditer(section_pattern, full_text))
        sections = []

        for idx, match in enumerate(matches):
            sec_type = match.group(1)
            sec_title = match.group(2).strip()
            level = 1 if sec_type == "section" else 2 if sec_type == "subsection" else 3

            start_pos = match.end()
            if idx + 1 < len(matches):
                end_pos = matches[idx + 1].start()
            else:
                end_doc = full_text.find("\\end{document}", start_pos)
                end_bib = full_text.find("\\begin{thebibliography}", start_pos)
                valid_ends = [p for p in [end_doc, end_bib] if p != -1]
                end_pos = min(valid_ends) if valid_ends else len(full_text)

            raw_content = full_text[start_pos:end_pos].strip()
            # Clean common LaTeX syntax from content for readability
            clean_content = re.sub(r"\\(textbf|textit|emph)\{([^}]+)\}", r"\2", raw_content)
            clean_content = re.sub(r"\\cite\{([^}]+)\}", r"[\1]", clean_content)
            clean_content = re.sub(r"\\ref\{([^}]+)\}", r"\1", clean_content)

            sections.append({
                "id": f"sec_{idx + 1}",
                "title": sec_title,
                "level": level,
                "content": clean_content.strip(),
                "order": idx + 1,
                "subsections": []
            })

        # Figures
        figures = []
        fig_matches = re.finditer(r"\\begin\{figure\}(.*?)\\end\{figure\}", full_text, re.DOTALL)
        for f_idx, fm in enumerate(fig_matches, 1):
            fig_block = fm.group(1)
            caption_m = re.search(r"\\caption\{([^}]+)\}", fig_block)
            label_m = re.search(r"\\label\{([^}]+)\}", fig_block)
            figures.append({
                "id": f"fig_{f_idx}",
                "caption": caption_m.group(1).strip() if caption_m else f"Figure {f_idx}",
                "label": label_m.group(1).strip() if label_m else f"fig:{f_idx}",
                "source": ""
            })

        # Tables
        tables = []
        tbl_matches = re.finditer(r"\\begin\{table\}(.*?)\\end\{table\}", full_text, re.DOTALL)
        for t_idx, tm in enumerate(tbl_matches, 1):
            tbl_block = tm.group(1)
            caption_m = re.search(r"\\caption\{([^}]+)\}", tbl_block)
            tables.append({
                "id": f"tbl_{t_idx}",
                "caption": caption_m.group(1).strip() if caption_m else f"Table {t_idx}",
                "data": []
            })

        return {
            "title": title,
            "abstract": abstract,
            "keywords": keywords,
            "authors": authors,
            "raw_text": full_text,
            "sections": sections,
            "figures": figures,
            "tables": tables
        }

    @classmethod
    def _structure_raw_text(cls, full_text: str) -> List[Dict[str, Any]]:
        headings_pattern = r"\n(?=[0-9IVX]+\.\s+|[A-Z][A-Za-z\s]{3,40}\n)"
        blocks = re.split(headings_pattern, full_text)

        sections = []
        sec_id = 1
        for block in blocks:
            clean_block = block.strip()
            if not clean_block:
                continue

            lines = clean_block.split("\n")
            sec_title = lines[0].strip()
            sec_content = "\n".join(lines[1:]).strip() if len(lines) > 1 else clean_block

            sections.append({
                "id": f"sec_{sec_id}",
                "title": sec_title[:100],
                "level": 1,
                "content": sec_content,
                "order": sec_id,
                "subsections": []
            })
            sec_id += 1

        if not sections:
            sections = [{
                "id": "sec_1",
                "title": "Main Content",
                "level": 1,
                "content": full_text,
                "order": 1,
                "subsections": []
            }]

        return sections
